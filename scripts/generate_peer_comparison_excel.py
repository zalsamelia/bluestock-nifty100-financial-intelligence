"""
Peer Comparison Excel Report Generator — Sprint 3, Day 20.
Generates output/peer_comparison.xlsx with one sheet per peer group (11 sheets).
Each sheet contains 20 KPI columns + percentile rank columns.
Conditional fills: green (>= 75th pct), yellow (25th-75th pct), red (<= 25th pct).
Benchmark company row is highlighted gold.
Peer group median summary row is added at the bottom.
"""

import sys
import sqlite3
import warnings
from pathlib import Path
from typing import Optional, Dict, Any, List

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

DB_PATH = PROJECT_ROOT / "nifty100.db"
PEER_GROUPS_PATH = PROJECT_ROOT / "data" / "raw" / "peer_groups.xlsx"
OUTPUT_DIR = PROJECT_ROOT / "output"

# 20 KPI columns to display per sheet
KPI_COLUMNS = [
    "return_on_equity_pct", "roce_pct", "net_profit_margin_pct", "return_on_assets_pct",
    "operating_profit_margin_pct",
    "debt_to_equity", "interest_coverage", "net_debt_cr",
    "free_cash_flow_cr", "cfo_quality_score", "capex_intensity_pct", "fcf_conversion_pct",
    "revenue_cagr_5yr", "pat_cagr_5yr", "eps_cagr_5yr",
    "asset_turnover", "book_value_per_share", "earnings_per_share",
    "composite_quality_score", "capital_allocation_pattern",
]

# Percentile-ranked metrics (10 core peer metrics)
PERCENTILE_METRICS = [
    "return_on_equity_pct", "roce_pct", "net_profit_margin_pct", "debt_to_equity",
    "free_cash_flow_cr", "pat_cagr_5yr", "revenue_cagr_5yr", "eps_cagr_5yr",
    "interest_coverage", "asset_turnover",
]

# Fills
GREEN_FILL  = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED_FILL    = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
GOLD_FILL   = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
MEDIAN_FILL = PatternFill(start_color="DDEEFF", end_color="DDEEFF", fill_type="solid")
HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
NORMAL_FONT = Font(name="Calibri", size=10)
BOLD_FONT   = Font(name="Calibri", bold=True, size=10)


def _get_percentile_fill(pct_rank: Optional[float]) -> Optional[PatternFill]:
    """Return fill based on percentile rank thresholds."""
    if pct_rank is None or pd.isna(pct_rank):
        return None
    if pct_rank >= 0.75:
        return GREEN_FILL
    if pct_rank >= 0.25:
        return YELLOW_FILL
    return RED_FILL


def _thin_border() -> Border:
    s = Side(style="thin", color="D9D9D9")
    return Border(left=s, right=s, top=s, bottom=s)


def _write_peer_sheet(
    ws,
    group_df: pd.DataFrame,
    peer_name: str,
    benchmark_ids: List[str],
    peer_percentiles_df: pd.DataFrame,
) -> None:
    """
    Write a single peer group sheet with KPI columns, percentile rank columns,
    green/yellow/red fills based on percentile rank, gold benchmark row, and median row.
    """
    border = _thin_border()

    # Build merged display dataframe
    avail_kpi = [c for c in KPI_COLUMNS if c in group_df.columns]
    display_df = group_df[["company_id", "company_name"] + avail_kpi].copy()

    # Merge percentile ranks per metric
    for metric in PERCENTILE_METRICS:
        col_pct_name = f"{metric}_pct_rank"
        pct_vals = {}
        for _, pr in peer_percentiles_df[
            (peer_percentiles_df["peer_group_name"] == peer_name) &
            (peer_percentiles_df["metric"] == metric)
        ].iterrows():
            pct_vals[pr["company_id"]] = pr["percentile_rank"]
        display_df[col_pct_name] = display_df["company_id"].map(pct_vals)

    # Sort by composite_quality_score desc
    if "composite_quality_score" in display_df.columns:
        display_df = display_df.sort_values("composite_quality_score", ascending=False)
    display_df = display_df.reset_index(drop=True)

    all_cols = list(display_df.columns)
    pct_rank_cols = [c for c in all_cols if c.endswith("_pct_rank")]
    metric_cols_base = set(PERCENTILE_METRICS)

    # Write headers
    for col_idx, col_name in enumerate(all_cols, start=1):
        header_text = col_name.replace("_pct_rank", " Pct").replace("_", " ").title()
        cell = ws.cell(row=1, column=col_idx)
        cell.value = header_text
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(col_idx)].width = 14 if "pct" in col_name.lower() else 16

    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"

    # Write data rows
    for row_idx, (_, row) in enumerate(display_df.iterrows(), start=2):
        cid = row["company_id"]
        is_benchmark = cid in benchmark_ids

        for col_idx, col_name in enumerate(all_cols, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = row[col_name]
            cell.value = round(float(val), 2) if isinstance(val, float) and pd.notna(val) else val
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = BOLD_FONT if is_benchmark else NORMAL_FONT

            # Benchmark: gold row background
            if is_benchmark:
                cell.fill = GOLD_FILL
            elif col_name in pct_rank_cols:
                pct_val = row[col_name]
                if pd.notna(pct_val):
                    cell.fill = _get_percentile_fill(pct_val) or PatternFill()

    # Median summary row at bottom
    median_row = len(display_df) + 2
    numeric_cols = display_df.select_dtypes(include=[float, int]).columns.tolist()
    medians = display_df[numeric_cols].median()

    ws.cell(row=median_row, column=1).value = "PEER MEDIAN"
    ws.cell(row=median_row, column=1).font = BOLD_FONT
    ws.cell(row=median_row, column=1).fill = MEDIAN_FILL
    ws.cell(row=median_row, column=1).border = border

    for col_idx, col_name in enumerate(all_cols, start=1):
        cell = ws.cell(row=median_row, column=col_idx)
        cell.border = border
        cell.fill = MEDIAN_FILL
        cell.font = BOLD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if col_name in medians.index and pd.notna(medians[col_name]):
            cell.value = round(float(medians[col_name]), 2)


def generate_peer_comparison_excel(
    output_path: Optional[Path] = None,
    target_year: int = 2024,
) -> Path:
    """
    Generate output/peer_comparison.xlsx with 11 peer group sheets.
    Each sheet has 20 KPI columns + percentile rank columns with full formatting.
    """
    output_path = output_path or (OUTPUT_DIR / "peer_comparison.xlsx")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    conn = sqlite3.connect(str(DB_PATH))

    # Load financials ratios
    ratios_df = pd.read_sql_query("""
        SELECT fr.*, c.company_name, s.sector, s.industry, s.market_cap_category
        FROM financial_ratios fr
        JOIN companies c ON fr.company_id = c.company_id
        LEFT JOIN sectors s ON fr.company_id = s.company_id
        WHERE fr.year = ?
    """, conn, params=(target_year,))

    # Load peer percentiles
    peer_pct_df = pd.read_sql_query("SELECT * FROM peer_percentiles", conn)
    conn.close()

    # Load peer group excel
    peer_map_df = pd.read_excel(PEER_GROUPS_PATH)
    peer_groups = sorted(peer_map_df["peer_group_name"].unique())

    # Get benchmark company IDs per peer group
    benchmark_map: Dict[str, List[str]] = {}
    for pg_name, pg_df in peer_map_df.groupby("peer_group_name"):
        benchmark_map[pg_name] = pg_df[pg_df["is_benchmark"] == True]["company_id"].tolist()

    # Compute composite quality score for the full set (for sorting)
    from src.screener.engine import calculate_composite_quality_score
    ratios_df = calculate_composite_quality_score(ratios_df)

    wb = Workbook()
    wb.remove(wb.active)

    for pg_name in peer_groups:
        pg_comps = peer_map_df[peer_map_df["peer_group_name"] == pg_name]["company_id"].tolist()
        pg_data = ratios_df[ratios_df["company_id"].isin(pg_comps)].copy()

        if pg_data.empty:
            continue

        benchmark_ids = benchmark_map.get(pg_name, [])
        sheet_title = pg_name[:31]
        ws = wb.create_sheet(title=sheet_title)
        _write_peer_sheet(ws, pg_data, pg_name, benchmark_ids, peer_pct_df)

    wb.save(str(output_path))
    print(f"Saved peer_comparison.xlsx: {output_path}")
    return output_path
